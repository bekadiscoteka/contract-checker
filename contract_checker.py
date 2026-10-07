#!/usr/bin/env python3
"""Сканер договоров: ищет рискованные пункты и строит отчёт.
Не является юридической консультацией.

Запуск:  python contract_checker.py договор.pdf [-o report.html] [--top 25]
Нужен ANTHROPIC_API_KEY. Без ключа работает режим правил (без LLM).
"""
import argparse, html, json, os, re, sys
from pathlib import Path

HERE = Path(__file__).parent
MODEL = os.getenv("CLAUDE_MODEL", "claude-sonnet-5-5")
LEVELS = {"none": 0, "low": 1, "medium": 2, "high": 3}
RU = {"none": "нет риска", "low": "низкий", "medium": "средний", "high": "высокий"}
DISCLAIMER = ("Отчёт создан автоматически и не является юридической консультацией. "
              "Он помогает найти потенциально рискованные формулировки для проверки юристом. "
              "Отсутствие находок не означает, что договор безопасен.")

# ---------- Perception: чтение и разбиение ----------
def load_text(path):
    p = Path(path); ext = p.suffix.lower()
    if ext == ".pdf":
        import pdfplumber
        with pdfplumber.open(p) as pdf:
            text = "\n".join((pg.extract_text() or "") for pg in pdf.pages)
        if len(text.strip()) < 100:  # вероятно скан: пробуем OCR
            try:
                from pdf2image import convert_from_path
                import pytesseract
                text = "\n".join(pytesseract.image_to_string(im, lang="rus+eng")
                                 for im in convert_from_path(str(p)))
            except Exception as e:
                sys.exit(f"Похоже, это скан, а OCR недоступен ({e}). Установите pytesseract и pdf2image.")
        return text
    if ext == ".docx":
        import docx
        d = docx.Document(str(p))
        parts = [x.text for x in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" ".join(c.text for c in row.cells))
        return "\n".join(parts)
    return p.read_text(encoding="utf-8", errors="ignore")

def split_clauses(text):
    text = re.sub(r"[ \t]+", " ", text)
    parts = re.split(r"\n(?=\s*\d+(?:\.\d+)*[.)]?\s)", text)
    if len(parts) < 3:
        parts = re.split(r"\n\s*\n", text)
    out = [re.sub(r"\s*\n\s*", " ", s).strip() for s in parts]
    return [s for s in out if len(s) > 30]

# ---------- Memory: база паттернов ----------
def load_patterns():
    return json.loads((HERE / "risk_patterns.json").read_text(encoding="utf-8"))

# ---------- Attention: ключевые слова + эмбеддинги ----------
def get_embedder():
    try:
        from sentence_transformers import SentenceTransformer
        return SentenceTransformer("paraphrase-multilingual-MiniLM-L12-v2")
    except Exception:
        print("sentence-transformers не найден: используются только ключевые слова.", file=sys.stderr)
        return None

def select_clauses(clauses, patterns, embedder, threshold=0.55):
    sims = None
    if embedder is not None:
        import numpy as np
        ex = [(i, e) for i, p in enumerate(patterns) for e in p["examples"]]
        E = embedder.encode([e for _, e in ex], normalize_embeddings=True)
        C = embedder.encode(clauses, normalize_embeddings=True)
        S = C @ E.T
        sims = np.zeros((len(clauses), len(patterns)))
        for col, (pi, _) in enumerate(ex):
            sims[:, pi] = np.maximum(sims[:, pi], S[:, col])
    selected = []
    for ci, c in enumerate(clauses):
        low = c.lower(); cats = []; score = 0.0
        for pi, p in enumerate(patterns):
            kw = any(k in low for k in p["keywords"])
            sm = float(sims[ci, pi]) if sims is not None else 0.0
            if kw or sm >= threshold:
                cats.append(p["category"])
                score = max(score, p["weight"] * (1 + sm + (0.5 if kw else 0)))
        if cats:
            selected.append({"id": ci + 1, "text": c, "categories": cats, "score": score,
                             "weight": max(p["weight"] for p in patterns if p["category"] in cats)})
    return sorted(selected, key=lambda x: -x["score"])

# ---------- Decision: анализ через Claude (zero-shot) ----------
SYSTEM = ("Ты помощник по проверке договоров. Для каждого пункта оцени риск для стороны, которая "
          "собирается подписать договор. Отвечай только JSON-массивом без пояснений и без markdown. "
          "Формат элемента: {\"id\": число, \"risk\": \"none|low|medium|high\", "
          "\"issue\": \"в чём проблема, 1-2 предложения\", \"recommendation\": \"что уточнить или изменить\"}. "
          "Пиши по-русски. Не давай юридических заключений и не пиши, что договор можно подписывать.")

def analyze_llm(items, batch=10):
    import anthropic
    client = anthropic.Anthropic()
    results = {}
    for i in range(0, len(items), batch):
        chunk = items[i:i + batch]
        prompt = "Проанализируй пункты договора:\n" + json.dumps(
            [{"id": x["id"], "categories": x["categories"], "text": x["text"]} for x in chunk],
            ensure_ascii=False)
        r = client.messages.create(model=MODEL, max_tokens=4000, system=SYSTEM,
                                   messages=[{"role": "user", "content": prompt}])
        raw = re.sub(r"^```(?:json)?|```$", "", r.content[0].text.strip(), flags=re.M).strip()
        try:
            for a in json.loads(raw):
                results[a["id"]] = a
        except Exception:
            print(f"Не удалось разобрать ответ для пачки {i // batch + 1}", file=sys.stderr)
    return results

def analyze_rules(items):
    lv = {1: "low", 2: "medium", 3: "high"}
    return {x["id"]: {"id": x["id"], "risk": lv[x["weight"]],
                      "issue": "Найдена типовая рискованная формулировка: " + ", ".join(x["categories"]) + ".",
                      "recommendation": "Проверьте условие внимательно и обсудите его с юристом."} for x in items}

# ---------- Aggregation ----------
def aggregate(items, analysis):
    pts = {"low": 1, "medium": 2, "high": 4}
    score = sum(pts.get(analysis.get(x["id"], {}).get("risk"), 0) for x in items)
    highs = sum(1 for x in items if analysis.get(x["id"], {}).get("risk") == "high")
    if highs >= 2 or score >= 10: level = "high"
    elif highs == 1 or score >= 4: level = "medium"
    elif score > 0: level = "low"
    else: level = "none"
    return level, score, highs

# ---------- Output: HTML-отчёт ----------
def build_report(name, total, items, analysis, level, score, mode):
    colors = {"high": "#b3261e", "medium": "#b26a00", "low": "#2e7d32", "none": "#5d6b7a"}
    rows = []
    for x in sorted(items, key=lambda x: -LEVELS.get(analysis.get(x["id"], {}).get("risk", "none"), 0)):
        a = analysis.get(x["id"], {}); r = a.get("risk", "none")
        if r == "none": continue
        rows.append(f'<div class="it" style="border-left-color:{colors[r]}"><h3>Пункт {x["id"]}: '
                    f'{html.escape(", ".join(x["categories"]))} — риск {RU[r]}</h3>'
                    f'<p class="q">{html.escape(x["text"])}</p>'
                    f'<p><b>Проблема:</b> {html.escape(a.get("issue", ""))}</p>'
                    f'<p><b>Что сделать:</b> {html.escape(a.get("recommendation", ""))}</p></div>')
    verdict = {"high": "Высокий риск. Обязательно покажите договор юристу до подписания.",
               "medium": "Средний риск. Рекомендуется проверка юристом.",
               "low": "Низкий риск, но отмеченные пункты стоит уточнить.",
               "none": "Явных красных флагов не найдено. Это не гарантия безопасности."}[level]
    return f"""<!DOCTYPE html><html lang="ru"><head><meta charset="utf-8"><title>Отчёт: {html.escape(name)}</title>
<style>body{{font:16px/1.5 Georgia,serif;max-width:800px;margin:24px auto;padding:0 16px;color:#1c2530}}
.sum{{border:1px solid #d5dce3;border-radius:6px;padding:14px 18px;margin-bottom:18px}}
.it{{border:1px solid #d5dce3;border-left:5px solid;border-radius:6px;padding:4px 16px;margin-bottom:12px}}
.it h3{{font:700 1rem system-ui,sans-serif;margin:10px 0 4px}}.q{{color:#5d6b7a;font-style:italic}}
.w{{background:#fff4e0;border-left:4px solid #b26a00;padding:8px 12px;font-size:.9rem}}</style></head><body>
<h1>Отчёт по договору: {html.escape(name)}</h1><p class="w">{DISCLAIMER}</p>
<div class="sum"><h2 style="color:{colors[level]};margin-top:0">{verdict}</h2>
<p>Пунктов в документе: {total}. Отобрано для анализа: {len(items)}. Баллы риска: {score}.<br>
Режим анализа: {mode}.</p></div>{"".join(rows) or "<p>Рискованных пунктов не выявлено.</p>"}
</body></html>"""

def main():
    ap = argparse.ArgumentParser(description="Сканер договоров (не юридическая консультация)")
    ap.add_argument("file"); ap.add_argument("-o", "--output", default="report.html")
    ap.add_argument("--top", type=int, default=25, help="максимум пунктов для анализа")
    ap.add_argument("--no-llm", action="store_true", help="только правила, без Claude")
    a = ap.parse_args()

    text = load_text(a.file)
    clauses = split_clauses(text)
    if not clauses: sys.exit("Не удалось выделить пункты договора.")
    items = select_clauses(clauses, load_patterns(), get_embedder())[:a.top]

    use_llm = not a.no_llm and os.getenv("ANTHROPIC_API_KEY")
    if use_llm:
        analysis = analyze_llm(items); mode = f"Claude ({MODEL}) + поиск по базе паттернов"
    else:
        print("Режим без LLM (нет ключа или указан --no-llm).", file=sys.stderr)
        analysis = analyze_rules(items); mode = "правила и ключевые слова, без LLM"
    level, score, highs = aggregate(items, analysis)

    Path(a.output).write_text(build_report(Path(a.file).name, len(clauses), items, analysis, level, score, mode),
                              encoding="utf-8")
    print(f"Итог: {RU[level]} риск, баллы {score}, высоких рисков {highs}. Отчёт: {a.output}")
    print(DISCLAIMER)

if __name__ == "__main__":
    main()
